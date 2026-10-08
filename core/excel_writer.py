"""
Writes the output workbook: Raw Data, Classified, Common Size, and Analysis
(incl. IRL) sheets. Ebrima size 10 is applied to every cell in every sheet,
per the "Additional Standing Rules" in the instructions doc.
"""

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

FONT_NAME = "Ebrima"
FONT_SIZE = 10
HEADER_FILL = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
SECTION_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")


def _header_font():
    return Font(name=FONT_NAME, size=FONT_SIZE, bold=True)


def _body_font():
    return Font(name=FONT_NAME, size=FONT_SIZE)


def _section_font():
    return Font(name=FONT_NAME, size=FONT_SIZE, bold=True, color="FFFFFF")


def _write_header_row(ws, headers, row=1):
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = _header_font()
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _autosize(ws, min_width=10, max_width=60):
    for column_cells in ws.columns:
        length = max((len(str(c.value)) for c in column_cells if c.value is not None), default=0)
        col_letter = column_cells[0].column_letter
        ws.column_dimensions[col_letter].width = max(min_width, min(max_width, length + 2))


def _write_raw_data(wb, line_items, fy_columns):
    ws = wb.create_sheet("Raw Data")
    headers = ["Statement", "Line Item", "Statement Tag"] + fy_columns
    _write_header_row(ws, headers)
    for row_idx, li in enumerate(line_items, start=2):
        ws.cell(row=row_idx, column=1, value=li.statement)
        ws.cell(row=row_idx, column=2, value=li.line_item)
        ws.cell(row=row_idx, column=3, value=li.statement_tag)
        for col_offset, fy in enumerate(fy_columns):
            ws.cell(row=row_idx, column=4 + col_offset, value=li.fy_values.get(fy))
    _autosize(ws)
    return ws


def _write_classified(wb, line_items, fy_columns):
    ws = wb.create_sheet("Classified")
    headers = ["Statement", "Line Item", "Category", "Confidence", "Matched Keywords"] + fy_columns
    _write_header_row(ws, headers)
    pnl_items = [li for li in line_items if li.statement == "P&L"]
    for row_idx, li in enumerate(pnl_items, start=2):
        ws.cell(row=row_idx, column=1, value=li.statement)
        ws.cell(row=row_idx, column=2, value=li.line_item)
        ws.cell(row=row_idx, column=3, value=li.category or li.statement_tag)
        ws.cell(row=row_idx, column=4, value=li.confidence)
        ws.cell(row=row_idx, column=5, value=", ".join(li.matched_keywords))
        for col_offset, fy in enumerate(fy_columns):
            ws.cell(row=row_idx, column=6 + col_offset, value=li.fy_values.get(fy))
    _autosize(ws)
    return ws


def _write_common_size(wb, line_items, fy_columns):
    ws = wb.create_sheet("Common Size")
    cs_headers = [f"{fy} %" for fy in fy_columns]
    yoy_headers = [f"{fy} YoY %" for fy in fy_columns[1:]]
    headers = ["Statement", "Line Item", "Category"] + cs_headers + yoy_headers
    _write_header_row(ws, headers)
    for row_idx, li in enumerate(line_items, start=2):
        ws.cell(row=row_idx, column=1, value=li.statement)
        ws.cell(row=row_idx, column=2, value=li.line_item)
        ws.cell(row=row_idx, column=3, value=li.category or li.statement_tag)
        col = 4
        for fy in fy_columns:
            ws.cell(row=row_idx, column=col, value=li.common_size.get(fy))
            col += 1
        for fy in fy_columns[1:]:
            ws.cell(row=row_idx, column=col, value=li.yoy_change_pct.get(fy))
            col += 1
    _autosize(ws)
    return ws


def _write_section(ws, row, title, findings):
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = _section_font()
    cell.fill = SECTION_FILL
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    row += 1
    if not findings:
        ws.cell(row=row, column=1, value="No findings for this section.")
        return row + 2
    for finding in findings:
        ws.cell(row=row, column=1, value=finding.area)
        ws.cell(row=row, column=2, value=finding.observation)
        ws.cell(row=row, column=3, value=finding.reason_for_concern)
        ws.cell(row=row, column=4, value=finding.info_required)
        row += 1
    return row + 1


def _write_analysis(wb, sections, irl_rows):
    ws = wb.create_sheet("Analysis")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 40

    row = 1
    for title in [
        "Executive Summary",
        "YoY Variance Analysis",
        "Common Size Analysis",
        "Ratio Commentary",
        "Data Integrity",
    ]:
        row = _write_section(ws, row, title, sections.get(title, []))

    # Information Requirement List
    irl_title_cell = ws.cell(row=row, column=1, value="Information Requirement List (IRL)")
    irl_title_cell.font = _section_font()
    irl_title_cell.fill = SECTION_FILL
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    row += 1

    irl_headers = [
        "Reference Number",
        "Financial Statement Area",
        "Observation",
        "Reason for Concern",
        "Information / Clarification Required",
    ]
    _write_header_row(ws, irl_headers, row=row)
    row += 1
    for irl_row in irl_rows:
        for col_idx, key in enumerate(irl_headers, start=1):
            ws.cell(row=row, column=col_idx, value=irl_row[key])
        row += 1

    ws.column_dimensions["E"].width = 40
    return ws


def _apply_font_everywhere(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                existing = cell.font
                cell.font = Font(
                    name=FONT_NAME,
                    size=FONT_SIZE,
                    bold=existing.bold,
                    color=existing.color,
                )


def write_workbook(output_path, company_name, line_items, fy_columns, sections, irl_rows):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    _write_raw_data(wb, line_items, fy_columns)
    _write_classified(wb, line_items, fy_columns)
    _write_common_size(wb, line_items, fy_columns)
    _write_analysis(wb, sections, irl_rows)

    # A simple cover/title row on each data sheet identifying the company.
    for sheet_name in ("Raw Data", "Classified", "Common Size", "Analysis"):
        ws = wb[sheet_name]
        ws.insert_rows(1)
        title_cell = ws.cell(row=1, column=1, value=f"{company_name} -- {sheet_name}")
        title_cell.font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True)

    _apply_font_everywhere(wb)
    wb.save(output_path)
    return output_path
