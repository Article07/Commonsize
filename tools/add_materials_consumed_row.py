#!/usr/bin/env python
"""
One-off change to the shipped template (template/commonsize_template.xlsx): a "Cost of materials consumed"
input row directly below "Purchase of traded goods" in the Sch. Cost of goods sold block, counted in the
block's total and in the IS Purchases line.

No row below moves (every IS / BS / CF / Ratios reference stays valid): the block's title and "Particulars"
heading move up into the blank separator row above them.

    before                                   after
    264  (blank)                             264  Cost of goods sold
    265  Cost of goods sold                  265  Particulars
    266  Particulars                         266  Purchase of traded goods
    267  Purchase of traded goods            267  Cost of materials consumed     (new)
    268  Change in inventories               268  Change in inventories
    271  Total = C267+C268                    271  Total = SUM(C266:C267)+C268
    IS 10  =Sch.!C267/curr                   IS 10  =SUM(Sch.!C266:C267)/curr

config/schedule_layout.json is updated to match (COGS-Purchase rows 266-267, title 264).

    python tools/add_materials_consumed_row.py
"""

import json
from copy import copy
from pathlib import Path

import openpyxl

APP_DIR = Path(__file__).resolve().parent.parent
TEMPLATE = APP_DIR / "template" / "commonsize_template.xlsx"
LAYOUT = APP_DIR / "config" / "schedule_layout.json"
COLS = range(1, 19)            # A..R: labels, year columns, % and y-o-y columns


def _copy_row(ws, src, dst):
    for c in COLS:
        a, b = ws.cell(src, c), ws.cell(dst, c)
        b.value = a.value
        if a.has_style:
            b._style = copy(a._style)
    ws.row_dimensions[dst].height = ws.row_dimensions[src].height


def main():
    wb = openpyxl.load_workbook(TEMPLATE)
    ws, is_ = wb["Sch."], wb["IS"]
    if str(ws["B267"].value).strip().lower().startswith("cost of materials"):
        print("Already applied.")
        return
    assert str(ws["B267"].value).strip() == "Purchase of traded goods", ws["B267"].value
    assert ws["B264"].value is None and str(ws["B265"].value).strip() == "Cost of goods sold"

    _copy_row(ws, 265, 264)     # title
    _copy_row(ws, 266, 265)     # Particulars + year headings
    _copy_row(ws, 267, 266)     # Purchase of traded goods
    _copy_row(ws, 266, 267)     # same input style for the new row
    ws["B267"].value = "Cost of materials consumed"
    for col in "CDEFG":
        ws[f"{col}267"].value = None
        ws[f"{col}271"].value = f"=SUM({col}266:{col}267)+{col}268"
        is_[f"{col}10"].value = f"=SUM(Sch.!{col}266:{col}267)/curr"
    is_["B10"].value = "Purchases"
    wb.save(TEMPLATE)

    layout = json.loads(LAYOUT.read_text(encoding="utf-8"))
    sch = layout["schedules"] if "schedules" in layout else layout
    sch["COGS-Purchase"].update(title_row=264, first_item_row=266, last_item_row=267)
    for key in ("Opening Stock (Change in Inventories)", "Closing Stock (Change in Inventories)"):
        sch[key]["title_row"] = 264
    LAYOUT.write_text(json.dumps(layout, indent=2) + "\n", encoding="utf-8")
    print("Template and layout updated.")


if __name__ == "__main__":
    main()
