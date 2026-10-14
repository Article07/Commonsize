#!/usr/bin/env python
"""
One-off change to the shipped template (template/commonsize_template.xlsx): an "Other reserves" balance row
(capital redemption reserve, general reserve ...) below "Securities premium" in the Sch. Reserves and surplus
block, counted in the block's total. No row below moves: the title and "Particulars" move up into the blank
row above them.

    before                                   after
    12  (blank)                              12  Reserves and surplus
    13  Reserves and surplus                 13  Particulars
    14  Particulars                          14  Securities premium
    15  Securities premium                   15  Other reserves                  (new)
    16  Surplus ...                          16  Surplus ...
    22  Total = C15+C21                      22  Total = C14+C15+C21

config/schedule_layout.json gets the new single-cell schedule "Other Reserves" (row 15) and Securities Premium
moves to row 14.

    python tools/add_other_reserves_row.py
"""

import json
from copy import copy
from pathlib import Path

import openpyxl

APP_DIR = Path(__file__).resolve().parent.parent
TEMPLATE = APP_DIR / "template" / "commonsize_template.xlsx"
LAYOUT = APP_DIR / "config" / "schedule_layout.json"
COLS = range(1, 19)


def _copy_row(ws, src, dst):
    for c in COLS:
        a, b = ws.cell(src, c), ws.cell(dst, c)
        b.value = a.value
        if a.has_style:
            b._style = copy(a._style)
    ws.row_dimensions[dst].height = ws.row_dimensions[src].height


def main():
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb["Sch."]
    if str(ws["B15"].value).strip().lower() == "other reserves":
        print("Already applied.")
        return
    assert ws["B12"].value is None and str(ws["B13"].value).strip() == "Reserves and surplus", ws["B13"].value
    assert str(ws["B15"].value).strip().lower() == "securities premium", ws["B15"].value

    _copy_row(ws, 13, 12)       # title
    _copy_row(ws, 14, 13)       # Particulars + year headings
    _copy_row(ws, 15, 14)       # Securities premium
    _copy_row(ws, 15, 15)       # same input style for the new row
    ws["B15"].value = "Other reserves"
    for col in "CDEFG":
        ws[f"{col}15"].value = None
        ws[f"{col}22"].value = f"={col}14+{col}15+{col}21"
    wb.save(TEMPLATE)

    layout = json.loads(LAYOUT.read_text(encoding="utf-8"))
    sch = layout["schedules"] if "schedules" in layout else layout
    for key in ("Reserves and Surplus", "Securities Premium", "Reserves Adjustments"):
        sch[key]["title_row"] = 12
    sch["Securities Premium"].update(first_item_row=14, last_item_row=14)
    sch["Other Reserves"] = dict(sch["Securities Premium"], first_item_row=15, last_item_row=15)
    if "order" in layout and "Other Reserves" not in layout["order"]:
        layout["order"].insert(layout["order"].index("Securities Premium") + 1, "Other Reserves")
    LAYOUT.write_text(json.dumps(layout, indent=2) + "\n", encoding="utf-8")
    print("Template and layout updated.")


if __name__ == "__main__":
    main()
