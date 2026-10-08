#!/usr/bin/env python
"""
Independent check of config/schedule_layout.json against the actual Reference
file, read with openpyxl (not the build tool's own bookkeeping): every
growable schedule's Total must sum exactly its item range, spare rows must be
blank + hidden, and item labels must look like what the layout claims.

    python tools/verify_layout.py
"""

import sys
from pathlib import Path

APP_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(APP_DIR))

import openpyxl

from config.schedule_map import ALL_SCHEDULE_ORDER, ALL_SCHEDULES, SHEET_NAME, spare_row_numbers

TEMPLATE = APP_DIR.parent / "Reference file" / "Common size format_Mfg-Trading firm.xlsx"


def main():
    wb = openpyxl.load_workbook(TEMPLATE)  # formulas, not cached values
    ws = wb[SHEET_NAME]
    errors = []
    print(f"{'schedule':38s} {'rows':>13s}  spare  first item / last item")
    for key in ALL_SCHEDULE_ORDER:
        s = ALL_SCHEDULES[key]
        first, last, total = s["first_item_row"], s["last_item_row"], s["total_row"]
        first_label = ws.cell(row=first, column=2).value
        last_label = ws.cell(row=last, column=2).value
        print(f"{key:38s} {first:>5d}-{last:<5d}{'':3s}{s['spare_rows']:>3d}   {first_label!r} / {last_label!r}")

        if s["pattern"] != "single_cell" and s["insertable"]:
            for col in "CDEFG":
                expected = f"=SUM({col}{first}:{col}{last})"
                actual = ws[f"{col}{total}"].value
                if actual != expected:
                    errors.append(f"{key}: {col}{total} is {actual!r}, expected {expected}")
        for r in spare_row_numbers(key):
            label = ws.cell(row=r, column=2).value
            hidden = ws.row_dimensions[r].hidden
            if label not in (None, "") or not hidden:
                errors.append(f"{key}: spare row {r} must be blank+hidden (label={label!r}, hidden={hidden})")
            for col in "CDEFG":
                if ws[f"{col}{r}"].value not in (None, ""):
                    errors.append(f"{key}: spare row {r} column {col} is not empty")
        if last_label in (None, "") and s["pattern"] != "single_cell":
            errors.append(f"{key}: last item row {last} has no label")

    print()
    if errors:
        print(f"{len(errors)} PROBLEM(S):")
        for e in errors:
            print("  -", e)
        sys.exit(1)
    print("Layout verified: every total spans its schedule, every spare row is blank and hidden.")


if __name__ == "__main__":
    main()
